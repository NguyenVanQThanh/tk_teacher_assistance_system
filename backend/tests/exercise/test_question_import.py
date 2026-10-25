# File Name: test_question_import.py
# Description: Integration tests for F-EXE-05 Question Import endpoint.
# Created At: 2026-08-09
# Created By: Tam Khang
# Updated At: 2026-08-09
# Updated By: Tam Khang

import io
import uuid
import docx
import pytest
from fastapi import status
from fastapi.testclient import TestClient


def get_admin_token(client: TestClient) -> str:
    resp = client.post(
        "/api/v1/auth/login",
        json={"phone": "0987654321", "password": "password123"}
    )
    return resp.json()["access_token"]


def create_sample_docx_with_answers() -> bytes:
    doc = docx.Document()
    doc.add_paragraph("PHẦN I. Câu hỏi trắc nghiệm 4 lựa chọn.")
    doc.add_paragraph("Câu 1. [NB] 2 + 2 x 3 bằng bao nhiêu?")
    doc.add_paragraph("A. 8")
    doc.add_paragraph("B. 12")
    doc.add_paragraph("C. 6")
    doc.add_paragraph("D. 10")

    doc.add_paragraph("PHẦN II. Câu hỏi trắc nghiệm đúng sai.")
    doc.add_paragraph("Câu 2. Cho phương trình x + 5 = 0.")
    doc.add_paragraph("a) Phương trình có nghiệm duy nhất.")
    doc.add_paragraph("b) Nghiệm của phương trình là x = 5.")

    doc.add_paragraph("PHẦN III. Câu hỏi trả lời ngắn.")
    doc.add_paragraph("Câu 3. [VD] Tìm nghiệm của x + 5 = 0.")

    table = doc.add_table(rows=4, cols=2)
    table.rows[0].cells[0].text = "Câu"
    table.rows[0].cells[1].text = "Đáp án"
    table.rows[1].cells[0].text = "1"
    table.rows[1].cells[1].text = "A"
    table.rows[2].cells[0].text = "2"
    table.rows[2].cells[1].text = "a) Đúng, b) Sai"
    table.rows[3].cells[0].text = "3"
    table.rows[3].cells[1].text = "-5"

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestQuestionImportEndpoint:
    def test_import_unsupported_file_type(self, client: TestClient):
        admin_token = get_admin_token(client)
        admin_headers = {"Authorization": f"Bearer {admin_token}"}

        # Create Subject & Exercise
        sub_resp = client.post(
            "/api/v1/subjects",
            json={"subject_name": f"Ly {uuid.uuid4().hex[:6]}"},
            headers=admin_headers
        )
        subject_id = sub_resp.json()["subject_id"]

        ex_resp = client.post(
            "/api/v1/exercises",
            json={"title": "De Vat Ly Test Import", "subject_id": subject_id},
            headers=admin_headers
        )
        exercise_id = ex_resp.json()["exercise_id"]

        # Upload .txt file -> 415
        resp = client.post(
            f"/api/v1/exercises/{exercise_id}/questions/import",
            files={"question_file": ("test.txt", b"some text", "text/plain")},
            headers=admin_headers
        )
        assert resp.status_code == status.HTTP_415_UNSUPPORTED_MEDIA_TYPE
        assert resp.json()["detail"]["code"] == "ERR-EXE-IMPORT_FILE_TYPE_UNSUPPORTED"

    def test_import_valid_docx(self, client: TestClient):
        admin_token = get_admin_token(client)
        admin_headers = {"Authorization": f"Bearer {admin_token}"}

        # Create Subject & Exercise
        sub_resp = client.post(
            "/api/v1/subjects",
            json={"subject_name": f"Sinh {uuid.uuid4().hex[:6]}"},
            headers=admin_headers
        )
        subject_id = sub_resp.json()["subject_id"]

        ex_resp = client.post(
            "/api/v1/exercises",
            json={"title": "De Sinh Test Import Docx", "subject_id": subject_id},
            headers=admin_headers
        )
        exercise_id = ex_resp.json()["exercise_id"]

        docx_bytes = create_sample_docx_with_answers()

        # Upload valid docx
        resp = client.post(
            f"/api/v1/exercises/{exercise_id}/questions/import",
            files={"question_file": ("de_thi_sample.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            headers=admin_headers
        )
        assert resp.status_code == status.HTTP_201_CREATED
        q_list = resp.json()
        assert len(q_list) == 3

        # Check total_questions in exercise detail
        ex_detail = client.get(f"/api/v1/exercises/{exercise_id}", headers=admin_headers)
        assert ex_detail.status_code == status.HTTP_200_OK
        assert ex_detail.json()["total_questions"] == 3
