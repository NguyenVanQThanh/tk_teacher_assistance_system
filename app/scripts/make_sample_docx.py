# File Name: make_sample_docx.py
# Description: Helper script to generate sample docx file with native Word Math Equations (OMML) and export to PDF.
# Created At: 2026-08-09
# Created By: Tam Khang
# Updated At: 2026-08-09
# Updated By: Tam Khang

import os
import re
import docx
import lxml.etree as ET
from latex2mathml.converter import convert
import win32com.client

# Locate MML2OMML.XSL stylesheet for MathML -> OMML transformation
XSL_PATH = "C:/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL"
if not os.path.exists(XSL_PATH):
    possible_paths = [
        "C:/Program Files (x86)/Microsoft Office/root/Office16/MML2OMML.XSL",
        "C:/Program Files/Microsoft Office/Office16/MML2OMML.XSL"
    ]
    for p in possible_paths:
        if os.path.exists(p):
            XSL_PATH = p
            break

xslt_transform = None
if os.path.exists(XSL_PATH):
    xslt_tree = ET.parse(XSL_PATH)
    xslt_transform = ET.XSLT(xslt_tree)


def latex_to_omml(latex_str: str):
    """Convert LaTeX string (e.g. '$Oxyz$') to native Word OMath XML element."""
    if not xslt_transform:
        return None
    try:
        clean_latex = latex_str.strip("$").strip()
        # Fix common LaTeX aliases
        clean_latex = clean_latex.replace("\\text{ cm}", "\\text{cm}")
        mathml_str = convert(clean_latex)
        mml_tree = ET.fromstring(mathml_str)
        omml_tree = xslt_transform(mml_tree)
        return omml_tree.getroot()
    except Exception as e:
        print(f"Warning: could not convert LaTeX '{latex_str}': {e}")
        return None


def add_paragraph_with_latex(doc, text: str, style=None):
    """Add paragraph to docx, replacing inline $...$ LaTeX with native Word OMath equations."""
    p = doc.add_paragraph(style=style)
    parts = re.split(r"(\$.*?\$)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("$") and part.endswith("$") and len(part) > 2:
            omml_elem = latex_to_omml(part)
            if omml_elem is not None:
                p._p.append(omml_elem)
            else:
                p.add_run(part)
        else:
            p.add_run(part)
    return p


def build_sample_docx(docx_path: str):
    doc = docx.Document()
    doc.add_heading("ĐỀ THI MINH HỌA (NATIVE WORD MATH EQUATIONS) - BỘ GD&ĐT", level=1)

    doc.add_paragraph("PHẦN I. Câu hỏi trắc nghiệm 4 lựa chọn.")
    add_paragraph_with_latex(doc, "Câu 1. [NB] Trong không gian $Oxyz$, cho điểm $A(1; 2; 3)$. Tọa độ hình chiếu vuông góc của $A$ trên mặt phẳng $(Oxy)$ là:")
    add_paragraph_with_latex(doc, "A. $(1; 2; 0)$    B. $(1; 0; 3)$    C. $(0; 2; 3)$    D. $(1; 2; -3)$")

    add_paragraph_with_latex(doc, "Câu 2. [TH] Tập nghiệm của bất phương trình $\\log_2(x - 1) < 3$ là:")
    add_paragraph_with_latex(doc, "A) $(1; 9)$    B) $(-\\infty; 9)$    C) $(1; +\\infty)$    D) $[1; 9]$")

    doc.add_paragraph("PHẦN II. Câu hỏi trắc nghiệm đúng sai.")
    add_paragraph_with_latex(doc, "Câu 3. Cho hàm số $y = f(x) = x^3 - 3x^2 + 2$.")
    add_paragraph_with_latex(doc, "a) Hàm số đạt cực đại tại $x = 0$.")
    add_paragraph_with_latex(doc, "b) Hàm số đồng biến trên khoảng $(0; 2)$.")
    add_paragraph_with_latex(doc, "c) Giá trị cực tiểu của hàm số bằng $-2$.")
    add_paragraph_with_latex(doc, "d) Đồ thị hàm số đi qua điểm $M(1; 0)$.")

    doc.add_paragraph("PHẦN III. Câu hỏi trả lời ngắn.")
    add_paragraph_with_latex(doc, "Câu 4. [VD] Một nhà sản xuất muốn thiết kế một chiếc thùng hình trụ có thể tích $V = 1000\\text{cm}^3$. Bán kính đáy $r$ (tính bằng cm) để diện tích toàn phần của thùng nhỏ nhất là bao nhiêu (làm tròn đến hàng phần mười)?")

    doc.add_heading("BẢNG ĐÁP ÁN MẪU (DẠNG 1 - NHIỀU MÃ ĐỀ)", level=2)
    tbl = doc.add_table(rows=5, cols=3)
    tbl.rows[0].cells[0].text = "Câu"
    tbl.rows[0].cells[1].text = "Mã đề 0121"
    tbl.rows[0].cells[2].text = "Mã đề 0122"

    tbl.rows[1].cells[0].text = "1"
    tbl.rows[1].cells[1].text = "A"
    tbl.rows[1].cells[2].text = "B"

    tbl.rows[2].cells[0].text = "2"
    tbl.rows[2].cells[1].text = "A"
    tbl.rows[2].cells[2].text = "C"

    tbl.rows[3].cells[0].text = "3"
    tbl.rows[3].cells[1].text = "a) Đúng\nb) Sai\nc) Đúng\nd) Đúng"
    tbl.rows[3].cells[2].text = "a) Sai\nb) Sai\nc) Đúng\nd) Sai"

    tbl.rows[4].cells[0].text = "4"
    tbl.rows[4].cells[1].text = "5,4"
    tbl.rows[4].cells[2].text = "5,4"

    doc.save(docx_path)
    print(f"Successfully generated {docx_path} with native Word Math Equations.")


def convert_docx_to_pdf(docx_path: str, pdf_path: str):
    abs_docx = os.path.abspath(docx_path)
    abs_pdf = os.path.abspath(pdf_path)
    print(f"Exporting PDF: {abs_docx} -> {abs_pdf}...")
    
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        doc = word.Documents.Open(abs_docx)
        doc.SaveAs(abs_pdf, FileFormat=17)  # 17 = wdFormatPDF
        doc.Close()
        print(f"Successfully exported PDF to: {abs_pdf}")
    finally:
        word.Quit()


if __name__ == "__main__":
    docx_file = "sample_template_bgd.docx"
    pdf_file = "sample_template_bgd.pdf"
    build_sample_docx(docx_file)
    convert_docx_to_pdf(docx_file, pdf_file)
